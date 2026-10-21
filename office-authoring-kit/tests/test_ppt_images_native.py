"""Regressions for omitted figures, risky geometry and false export success."""
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import Mock, patch
from zipfile import ZipFile
import xml.etree.ElementTree as E

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT/'powerpoint-authoring/scripts'))
sys.path.insert(0, str(ROOT/'word-authoring/scripts'))
import create_pptx
import check_pptx
import powerpoint_macos as native
import verify_office_macos as office
from pptx import Presentation
from pptx.util import Inches
from PIL import Image


class ImageAndNativeTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name)
        self.base = ROOT/'powerpoint-authoring/examples'
        self.spec = json.loads((self.base/'illustrated.json').read_text(encoding='utf-8'))

    def ppt(self, spec=None):
        path = self.path/'Editable 图文.pptx'
        create_pptx.build(spec or self.spec, self.base).save(path)
        return path

    def test_figures_icons_and_no_image_slide_follow_plan_and_preserve_bytes(self):
        p = self.ppt();before = p.read_bytes()
        r = check_pptx.check(p, self.spec, self.base)
        self.assertFalse(r['errors'])
        self.assertEqual([len(s['pictures']) for s in r['slides']], [1, 2, 0])
        expected = hashlib.sha256((self.base/'assets/synthetic-signal.png').read_bytes()).hexdigest()
        self.assertEqual(r['slides'][0]['pictures'][0]['sha256'], expected)
        self.assertIn('MIT', r['slides'][0]['notes'])
        self.assertEqual(before, p.read_bytes())

    def test_source_hyperlink_and_credit_survive_reopen(self):
        sp = deepcopy(self.spec)
        sp['assets']['signal']['source_url'] = 'https://example.org/original-figure'
        r = check_pptx.check(self.ppt(sp), sp, self.base)
        self.assertFalse(r['errors'])
        self.assertEqual(r['slides'][0]['pictures'][0]['source_url'], sp['assets']['signal']['source_url'])
        self.assertIn(sp['assets']['signal']['credit'], '\n'.join(r['slides'][0]['text']))

    def test_missing_file_credit_and_changed_asset_checksum_are_rejected(self):
        for key, value, error in [('path', 'absent.png', FileNotFoundError),
                                   ('credit', '', ValueError), ('sha256', '0'*64, ValueError)]:
            with self.subTest(key=key):
                sp = deepcopy(self.spec);sp['assets']['signal'][key] = value
                with self.assertRaises(error):create_pptx.build(sp, self.base)

    def test_downloaded_file_and_caption_do_not_mask_a_missing_embedded_image(self):
        p = self.ppt();prs = Presentation(p)
        slide = prs.slides[0]
        pic = next(s for s in slide.shapes if s.shape_type == 13)
        slide.shapes._spTree.remove(pic._element)
        prs.save(p)
        self.assertTrue(check_pptx.check(p, self.spec, self.base)['errors'])

    def test_crop_of_contain_only_planned_image_is_not_silently_accepted(self):
        p = self.ppt();prs = Presentation(p)
        next(s for s in prs.slides[0].shapes if s.shape_type == 13).crop_left = .1
        prs.save(p)
        self.assertTrue(check_pptx.check(p, self.spec, self.base)['errors'])

    def test_stretched_picture_is_reported(self):
        p = self.ppt();prs = Presentation(p)
        picture = next(s for s in prs.slides[0].shapes if s.shape_type == 13)
        picture.width = int(picture.width * .6);prs.save(p)
        self.assertIn('stretched_image', {w['kind'] for w in check_pptx.check(p)['warnings']})

    def test_inside_canvas_text_overlap_and_short_box_require_review(self):
        p = self.ppt();prs = Presentation(p);sl = prs.slides[2]
        create_pptx.text(sl, 'first\nsecond\nthird', 1, 2, 4, .2, 24, '000000', 'Arial')
        create_pptx.text(sl, 'neighbour', 1, 2.1, 4, .5, 24, '000000', 'Arial')
        prs.save(p);r = check_pptx.check(p)
        self.assertFalse(r['errors'])
        self.assertEqual(r['status'], 'review_required')
        self.assertTrue({'intersecting_text_regions','explicit_text_height'} <= {w['kind'] for w in r['warnings']})

    def test_plain_style_changes_native_colour_without_modifying_colour_image(self):
        p = self.ppt();ns = {'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
        with ZipFile(p) as z:
            colors = {c.get('val') for n in z.namelist() if n.startswith('ppt/slides/slide') and n.endswith('.xml')
                      for c in E.fromstring(z.read(n)).findall('.//a:srgbClr',ns)}
            self.assertFalse(colors-{'000000','FFFFFF'})
            theme = E.fromstring(z.read('ppt/theme/theme1.xml'))
            self.assertEqual(theme.find('.//a:hlink/a:srgbClr',ns).get('val'),'000000')
        image = next(s for s in Presentation(p).slides[0].shapes if s.shape_type == 13)
        self.assertEqual(image.image.blob,(self.base/'assets/synthetic-signal.png').read_bytes())

    @unittest.skipUnless(shutil.which('pdfinfo'), 'Poppler not installed')
    def test_partial_invalid_and_wrong_page_count_pdfs_do_not_pass(self):
        p=self.path/'proof.pdf';im=Image.new('RGB',(50,50),'white')
        im.save(p,'PDF',save_all=True,append_images=[im,im])
        self.assertTrue(native.complete_pdf(p,3))
        self.assertFalse(native.complete_pdf(p,2))
        data=p.read_bytes();p.write_bytes(data[:-20])
        self.assertFalse(native.complete_pdf(p,3))
        p.write_bytes(b'%PDF-1.4\ninvalid body\n%%EOF')
        self.assertFalse(native.complete_pdf(p,3))

    def test_already_open_target_is_preserved_without_launch_or_export(self):
        p=self.ppt();r={'source_sha256':hashlib.sha256(p.read_bytes()).hexdigest()}
        with patch.object(native.shutil,'which',return_value='/fake/pdfinfo'), patch.object(native,'apple',return_value='3|test'), patch.object(native.subprocess,'run') as run:
            with self.assertRaisesRegex(RuntimeError,'already open'):native.verify_ppt(p,self.path/'out.pdf',r)
            run.assert_not_called()
        self.assertFalse(r['opened_by_helper'])

    def test_success_reply_without_pdf_leaves_export_failed_and_target_open(self):
        p=self.ppt();folder=self.path/'export';folder.mkdir()
        r={'source_sha256':hashlib.sha256(p.read_bytes()).hexdigest()}
        proc=Mock();proc.communicate.return_value=('PPT_EXPORT_OK','');proc.returncode=0
        with patch.object(native.shutil,'which',return_value='/fake/pdfinfo'), \
             patch.object(native,'apple',side_effect=['NONE','3|test']) as apple, \
             patch.object(native.subprocess,'run'),patch.object(native,'start_export',return_value=proc), \
             patch.object(native.tempfile,'mkdtemp',return_value=str(folder)):
            with self.assertRaisesRegex(RuntimeError,'No stable, complete'):
                native.verify_ppt(p,self.path/'out.pdf',r,export_seconds=.01)
            self.assertEqual(apple.call_count,2)  # No close or dialog interaction.
        self.assertEqual(r['native_open'],'passed')
        self.assertEqual(r['native_export'],'failed')
        self.assertTrue(r['owned_target_left_open'])
        self.assertFalse((self.path/'out.pdf').exists())

    def test_existing_output_is_refused_before_opening_office(self):
        source=self.ppt();pdf=self.path/'keep.pdf';pdf.write_bytes(b'user content')
        with patch.object(office.platform,'system',return_value='Darwin'),patch.object(native,'verify_ppt') as verify:
            with self.assertRaises(FileExistsError):office.verify(source,pdf,self.path/'receipt.json')
            verify.assert_not_called()
        self.assertEqual(pdf.read_bytes(),b'user content')


if __name__ == '__main__':unittest.main()
