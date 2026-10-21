"""Behavior tests for file preservation, native object data and bad packages."""
import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "word-authoring/scripts"))
sys.path.insert(0, str(ROOT / "powerpoint-authoring/scripts"))
import create_docx
import create_pptx
import check_office_package
import check_pptx
from pptx.util import Inches
from docx import Document

spec = importlib.util.spec_from_file_location("owned_deploy", ROOT / "office-authoring-kit/deploy.py")
deploy = importlib.util.module_from_spec(spec)
spec.loader.exec_module(deploy)


class OfficeTests(unittest.TestCase):
    def test_deployment_plan_covers_group_members_without_writing(self):
        with tempfile.TemporaryDirectory() as temporary:
            destination = Path(temporary) / "discovery"
            declared = json.loads((ROOT / "GROUP.json").read_text(encoding="utf-8"))["skills"]
            plan = deploy.deploy(destination)
            self.assertEqual({item["name"] for item in plan["skills"]}, set(declared))
            self.assertFalse(destination.exists())
            self.assertFalse(plan["applied"])

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name)

    def word(self):
        spec = json.loads((ROOT / "word-authoring/examples/document.json").read_text(encoding="utf-8"))
        p = self.path / "A 中文 document.docx"
        create_docx.build(spec, self.path).save(p)
        return p

    def ppt(self):
        spec = json.loads((ROOT / "powerpoint-authoring/examples/presentation.json").read_text(encoding="utf-8"))
        p = self.path / "A 中文 presentation.pptx"
        create_pptx.build(spec).save(p)
        return p

    def mutate(self, source, name, transform):
        target = self.path / ("bad" + source.suffix)
        with zipfile.ZipFile(source) as src, zipfile.ZipFile(target, "w") as out:
            for item in src.infolist():
                data = src.read(item)
                out.writestr(item, transform(data) if item.filename == name else data)
        return target

    def test_bilingual_word_table_semantic_heading_and_page_field_survive(self):
        p = self.word()
        d = Document(p)
        self.assertIn("中文内容", "".join(x.text for x in d.paragraphs))
        self.assertEqual(d.tables[0].cell(2, 0).text, "Native opening")
        self.assertTrue(any(x.style.name == "Heading 1" for x in d.paragraphs))
        with zipfile.ZipFile(p) as z:
            self.assertIn(b'PAGE', z.read('word/footer1.xml'))
        before = p.read_bytes()
        self.assertEqual(check_office_package.check(p)["status"], "passed")
        self.assertEqual(before, p.read_bytes())

    def test_native_chart_data_and_notes_survive(self):
        p = self.ppt()
        result = check_pptx.check(p)
        self.assertEqual(result["slide_count"], 5)
        chart = result["slides"][3]["charts"][0]
        self.assertEqual(chart["categories"], ["A", "B", "C"])
        self.assertEqual(chart["series"][0]["values"], [12, 18, 24])
        self.assertIn("A=12", result["slides"][3]["notes"])
        self.assertEqual(check_office_package.check(p)["status"], "passed")

    def test_missing_namespace_in_pptx_is_detected(self):
        p = self.ppt()
        bad = self.mutate(p, "ppt/slides/slide1.xml", lambda b: b.replace(b'<p:sld ',
            b'<p:sld xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" mc:Ignorable="future" ', 1))
        result = check_office_package.check(bad)
        self.assertEqual(result["status"], "failed")
        self.assertEqual(result["issues"][0]["undefined_prefix"], "future")

    def test_broken_relationship_target_is_detected(self):
        p = self.word()
        bad = self.mutate(p, "_rels/.rels", lambda b: b.replace(b'word/document.xml', b'word/absent.xml'))
        self.assertTrue(check_office_package.check(bad)["package_errors"])

    def test_top_level_shape_outside_canvas_fails(self):
        p = self.ppt()
        from pptx import Presentation
        prs = Presentation(p)
        prs.slides[0].shapes[0].left = Inches(20)
        prs.save(p)
        self.assertEqual(check_pptx.check(p)["status"], "failed")

    def test_missing_image_and_unsupported_fields_fail(self):
        with self.assertRaises(FileNotFoundError):
            create_docx.build({"blocks": [{"type": "image", "path": "missing.png"}]}, self.path)
        with self.assertRaises(ValueError):
            create_pptx.build({"slides": [{"layout": "content", "title": "A", "body": "B", "hidden_data": "C"}]})

    def test_cli_refuses_existing_output_without_changing_it(self):
        for name, generator, example, suffix in [
            ("word-authoring", "create_docx.py", "document.json", ".docx"),
            ("powerpoint-authoring", "create_pptx.py", "presentation.json", ".pptx")]:
            output = self.path / ("keep" + suffix)
            output.write_bytes(b"existing user content")
            result = subprocess.run([sys.executable, str(ROOT/name/"scripts"/generator),
                str(ROOT/name/"examples"/example), str(output)], capture_output=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertEqual(output.read_bytes(), b"existing user content")

    def test_deployment_dry_run_idempotence_and_conflict_protection(self):
        dest = self.path / "skills"
        deploy.deploy(dest)
        self.assertFalse(dest.exists())
        conflict = self.path / "conflict"
        (conflict/"powerpoint-authoring").mkdir(parents=True)
        with self.assertRaises(FileExistsError):
            deploy.deploy(conflict, apply=True)
        self.assertFalse((conflict/"word-authoring").exists())

    def test_deployment_symlink_idempotence(self):
        dest = self.path / "skills"
        try:
            deploy.deploy(dest, apply=True)
        except OSError as exc:
            if getattr(exc, "winerror", None) == 1314:
                self.skipTest("Windows directory symlink creation privilege is unavailable")
            raise
        deploy.deploy(dest, apply=True)
        self.assertEqual((dest/"word-authoring").resolve(), ROOT/"word-authoring")
        self.assertEqual((dest/"excel-authoring").resolve(), ROOT/"excel-authoring")
        self.assertFalse((dest/"excel-live-control").exists())


if __name__ == "__main__":
    unittest.main()
