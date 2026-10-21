"""Regressions for small projected text, autofit and inherited template collisions."""
import json
from pathlib import Path
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT/'powerpoint-authoring/scripts'))
import create_pptx
import check_pptx
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt


class ReadabilityTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.path = Path(tmp.name)/'test.pptx'
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])

    def check(self, **kwargs):
        self.prs.save(self.path)
        before = self.path.read_bytes()
        result = check_pptx.check(self.path, **kwargs)
        self.assertEqual(self.path.read_bytes(), before)
        return result

    def test_small_run_and_table_are_flagged_despite_valid_bounds(self):
        shape = create_pptx.text(self.slide, 'Large ', 1, 1, 5, 1, 24, '000000', 'Arial')
        run = shape.text_frame.paragraphs[0].add_run()
        run.text = 'Small qualification'; run.font.size = Pt(12)
        table = self.slide.shapes.add_table(1, 1, Inches(1), Inches(3), Inches(5), Inches(1)).table
        p = table.cell(0, 0).text_frame.paragraphs[0]
        p.text = 'Small table result'; p.font.size = Pt(14)
        result = self.check(min_font_pt=22)
        self.assertFalse(result['errors'])
        small = [w for w in result['warnings'] if w['kind']=='small_font_review']
        self.assertEqual({w['location'] for w in small}, {'shape', 'cell 1,1'})
        self.assertEqual({n for w in small for n in w['font_sizes_pt']}, {12, 14})
        self.assertEqual(result['status'], 'review_required')

    def test_shrink_and_unknown_inheritance_cannot_pass_as_large_type(self):
        shape = create_pptx.text(self.slide, 'Nominal large text', 1, 1, 2, .3, 24, '000000', 'Arial')
        shape.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        other = self.slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
        other.text = 'Inherited type is unresolved'
        kinds = {w['kind'] for w in self.check(min_font_pt=22)['warnings']}
        self.assertTrue({'shrink_to_fit_review', 'inherited_font_size_unassessed'} <= kinds)

    def test_threshold_is_configurable_without_auto_changing_content(self):
        create_pptx.text(self.slide, 'Reading copy', 1, 1, 4, 1, 18, '000000', 'Arial')
        self.assertIn('small_font_review', {w['kind'] for w in self.check(min_font_pt=22)['warnings']})
        self.assertEqual(self.check(min_font_pt=18)['warnings'], [])

    def test_decorative_top_bar_can_hit_reserved_logo_with_no_text_collision(self):
        create_pptx.rect(self.slide, 0, 0, 13.333, .2, '000000')
        create_pptx.text(self.slide, 'Safe title', .5, 1, 6, 1, 30, '000000', 'Arial')
        # Simulate branding inherited from a template layout, absent from slide.shapes.
        logo = create_pptx.text(self.slide, 'LOGO', 11, 0, 2, .6, 24, '000000', 'Arial')
        # Text-only synthetic shape has no media/relationship dependencies.
        self.slide.slide_layout.shapes._spTree.append(logo._element)
        zones = [{'slide':'all', 'label':'inherited logo', 'rect':[11, 0, 2, .6]}]
        result = self.check(protected_regions=zones)
        self.assertFalse(result['errors'])
        overlaps = [w for w in result['warnings'] if w['kind']=='protected_region_intersection']
        self.assertEqual(len(overlaps), 1)
        self.assertEqual(overlaps[0]['shape'], self.slide.shapes[0].shape_id)
        self.assertEqual(overlaps[0]['region'], 'inherited logo')

    def test_regions_apply_only_to_targeted_slides(self):
        self.prs.slides.add_slide(self.prs.slide_layouts[6])
        create_pptx.rect(self.slide, 11, 0, 2, .5, '000000')
        zones = [{'slide':2, 'label':'second layout only', 'rect':[11, 0, 2, .6]}]
        self.assertEqual(self.check(protected_regions=zones)['warnings'], [])

    def test_bad_thresholds_and_regions_fail_explicitly(self):
        for value in (0, -1, float('nan'), True):
            with self.subTest(value=value), self.assertRaises(ValueError):
                self.check(min_font_pt=value)
        for value in ({}, [{'slide':2,'label':'x','rect':[1,1,1,1]}],
                      [{'slide':'all','label':'x','rect':[1,1,-1,1]}],
                      [{'slide':'all','label':'x','rect':[1,1,float('inf'),1]}],
                      [{'slide':'all','label':'x','rect':[13,1,1,1]}]):
            with self.subTest(value=value), self.assertRaises(ValueError):
                self.check(protected_regions=value)

    def test_chart_axes_and_labels_are_measured_and_defaults_are_projection_sized(self):
        spec = json.loads((ROOT/'powerpoint-authoring/examples/presentation.json').read_text(encoding='utf-8'))
        self.prs = create_pptx.build(spec)
        result = self.check(min_font_pt=22)
        chart_entry = next(s for s in result['slides'] if s['charts'])
        chart_fonts = [t for t in chart_entry['typography'] if t['location'] != 'shape']
        self.assertGreaterEqual(len(chart_fonts), 4)
        self.assertTrue(all(t['font_sizes_pt']==[22.0] for t in chart_fonts))
        self.prs.slides[3].shapes[2].chart.category_axis.tick_labels.font.size = Pt(12)
        result = self.check(min_font_pt=22)
        self.assertTrue(any(w['kind']=='small_font_review' and w['location']=='category_axis'
                            for w in result['warnings']))

    def test_explanatory_caption_is_projection_sized_and_autofit_disabled(self):
        base = ROOT/'powerpoint-authoring/examples'
        spec = json.loads((base/'illustrated.json').read_text(encoding='utf-8'))
        self.prs = create_pptx.build(spec, base)
        self.check(min_font_pt=22)
        caption = next(s for s in self.prs.slides[0].shapes
                       if s.has_text_frame and s.text==spec['slides'][0]['caption'])
        self.assertEqual(caption.text_frame.paragraphs[0].font.size.pt, 22)
        self.assertEqual(caption.text_frame.auto_size, MSO_AUTO_SIZE.NONE)


if __name__ == '__main__':
    unittest.main()
