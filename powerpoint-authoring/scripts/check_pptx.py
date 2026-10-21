#!/usr/bin/env python3
"""Inventory PPTX content/images and flag layout risks; not a visual-fit proof."""
import argparse
from collections import Counter
import hashlib
import io
import json
import math
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx_assets import load_asset, slide_asset_ids

EMU_PT = 12700
EMU_IN = 914400


def box(shape):
    return [shape.left / EMU_PT, shape.top / EMU_PT,
            (shape.left + shape.width) / EMU_PT, (shape.top + shape.height) / EMU_PT]


def text_height_risk(shape):
    """Conservative explicit-line lower bound, deliberately not a wrapping model."""
    frame = shape.text_frame
    required = 0
    paragraphs = list(frame.paragraphs)
    for index, p in enumerate(paragraphs):
        sizes = [r.font.size.pt for r in p.runs if r.font.size is not None]
        if p.font.size is not None:
            sizes.append(p.font.size.pt)
        if not sizes:  # Cannot infer inherited theme fonts reliably here.
            return None
        size = max(sizes)
        spacing = p.line_spacing
        if isinstance(spacing, float):
            line_height = size * spacing
        elif spacing is not None:
            line_height = spacing.pt
        else:
            line_height = size
        required += line_height * (p.text.count('\v') + 1)
        required += (p.space_before.pt if p.space_before is not None else 0)
        # Last-paragraph after-spacing need not occupy text height.
        if index < len(paragraphs) - 1:
            required += p.space_after.pt if p.space_after is not None else 0
    available = (shape.height - frame.margin_top - frame.margin_bottom) / EMU_PT
    if required > available + 2:
        return {'minimum_line_height_pt': round(required, 2), 'available_height_pt': round(available, 2)}
    return None


def font_inventory(frame):
    """Explicit run/paragraph sizes only; keep unresolved inheritance visible."""
    sizes, unresolved = set(), 0
    for paragraph in frame.paragraphs:
        if not paragraph.text.strip():
            continue
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            size = run.font.size if run.font.size is not None else paragraph.font.size
            if size is None:
                unresolved += 1
            else:
                sizes.add(round(size.pt, 2))
        if not paragraph.runs:
            unresolved += 1
    return sorted(sizes), unresolved


def validate_regions(regions, slide_count, width, height):
    if regions is None:
        return []
    if not isinstance(regions, list):
        raise ValueError('Protected regions must be a JSON array')
    for region in regions:
        if not isinstance(region, dict) or set(region) != {'slide', 'label', 'rect'}:
            raise ValueError('Each protected region needs only slide, label and rect')
        target = region['slide']
        if target != 'all' and (type(target) is not int or not 1 <= target <= slide_count):
            raise ValueError('Region slide must be all or a valid 1-based slide number')
        if not isinstance(region['label'], str) or not region['label'].strip():
            raise ValueError('Region label must be nonempty')
        rect = region['rect']
        if not isinstance(rect, list) or len(rect) != 4 or any(
                type(n) not in (float, int) or not math.isfinite(n) for n in rect):
            raise ValueError('Region rect must be four finite numbers in inches')
        x, y, w, h = rect
        if x < 0 or y < 0 or w <= 0 or h <= 0 or x+w > width+.001 or y+h > height+.001:
            raise ValueError('Protected region must have positive area inside the slide')
    return regions


def check(path, spec=None, base=None, min_font_pt=None, protected_regions=None):
    path = Path(path)
    payload = path.read_bytes()
    prs = Presentation(io.BytesIO(payload))
    if min_font_pt is not None and (type(min_font_pt) not in (float, int) or
            not math.isfinite(min_font_pt) or min_font_pt <= 0):
        raise ValueError('Minimum font size must be a positive finite number')
    regions = validate_regions(protected_regions, len(prs.slides),
                               prs.slide_width / EMU_IN, prs.slide_height / EMU_IN)
    errors, warnings, slides = [], [], []
    if spec is not None and len(spec.get('slides', [])) != len(prs.slides):
        errors.append('Plan/PPTX slide count mismatch')
    base = Path(base or '.')
    for i, slide in enumerate(prs.slides, 1):
        entry = {'slide': i, 'text': [], 'charts': [], 'pictures': [], 'notes': '', 'typography': []}
        text_regions = []

        def type_check(frame, shape_id, location):
            if not frame.text.strip():
                return
            auto = frame._txBody.bodyPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}normAutofit')
            if frame.auto_size == MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE or auto is not None:
                warnings.append({'slide': i, 'shape': shape_id, 'location': location,
                                 'kind': 'shrink_to_fit_review',
                                 'font_scale': auto.get('fontScale') if auto is not None else None})
            if min_font_pt is None:
                return
            sizes, unknown = font_inventory(frame)
            entry['typography'].append({'shape': shape_id, 'location': location,
                                        'font_sizes_pt': sizes, 'unresolved_runs': unknown})
            small = [size for size in sizes if size < min_font_pt]
            if small:
                warnings.append({'slide': i, 'shape': shape_id, 'location': location,
                                 'kind': 'small_font_review', 'font_sizes_pt': small,
                                 'threshold_pt': min_font_pt})
            if unknown:
                warnings.append({'slide': i, 'shape': shape_id, 'location': location,
                                 'kind': 'inherited_font_size_unassessed', 'unresolved_runs': unknown})

        def chart_font(font, shape_id, location):
            if min_font_pt is None:
                return
            size = font.size.pt if font.size is not None else None
            entry['typography'].append({'shape': shape_id, 'location': location,
                                        'font_sizes_pt': [size] if size is not None else [],
                                        'unresolved_runs': int(size is None)})
            if size is None or size < min_font_pt:
                warnings.append({'slide': i, 'shape': shape_id, 'location': location,
                                 'kind': 'inherited_font_size_unassessed' if size is None else 'small_font_review',
                                 'font_sizes_pt': [size] if size is not None else [], 'threshold_pt': min_font_pt})

        def inspect(s, grouped=False):
            if not grouped:
                sb = box(s)
                for region in regions:
                    if region['slide'] not in ('all', i):
                        continue
                    x, y, w, h = region['rect']
                    overlap_w = min(sb[2], (x+w)*72) - max(sb[0], x*72)
                    overlap_h = min(sb[3], (y+h)*72) - max(sb[1], y*72)
                    if overlap_w > .1 and overlap_h > .1:
                        warnings.append({'slide': i, 'shape': s.shape_id, 'kind': 'protected_region_intersection',
                                         'region': region['label'],
                                         'intersection_pt': [round(overlap_w, 2), round(overlap_h, 2)]})
            if not grouped and (s.left < 0 or s.top < 0 or s.width < 0 or s.height < 0 or
                    s.left + s.width > prs.slide_width + 20 or s.top + s.height > prs.slide_height + 20):
                errors.append(f'Slide {i}, shape {s.shape_id}: outside canvas')
            if s.shape_type == MSO_SHAPE_TYPE.GROUP:
                warnings.append({'slide': i, 'shape': s.shape_id, 'kind': 'group_geometry_unassessed'})
                for child in s.shapes:
                    inspect(child, True)
                return
            if s.has_text_frame:
                entry['text'].append(s.text)
                type_check(s.text_frame, s.shape_id, 'shape')
                if s.text.strip() and not grouped:
                    text_regions.append((s.shape_id, box(s)))
                    risk = text_height_risk(s)
                    if risk:
                        warnings.append(dict(slide=i, shape=s.shape_id, kind='explicit_text_height', **risk))
            if s.has_table:
                entry['text'].extend(' | '.join(c.text for c in row.cells) for row in s.table.rows)
                for row_number, row in enumerate(s.table.rows, 1):
                    for col_number, cell in enumerate(row.cells, 1):
                        type_check(cell.text_frame, s.shape_id, f'cell {row_number},{col_number}')
                if not grouped:
                    text_regions.append((s.shape_id, box(s)))
            if s.has_chart:
                chart = s.chart
                entry['charts'].append({'categories': [c.label for c in chart.plots[0].categories],
                    'series': [{'name': se.name, 'values': list(se.values)} for se in chart.series]})
                if min_font_pt is not None:
                    # Axis/plot defaults do not cover individual point or label overrides.
                    warnings.append({'slide': i, 'shape': s.shape_id, 'kind': 'chart_typography_partial',
                                     'detail': 'Per-point overrides and theme inheritance need native review'})
                    if chart.has_title:
                        type_check(chart.chart_title.text_frame, s.shape_id, 'chart title')
                    if chart.has_legend:
                        chart_font(chart.legend.font, s.shape_id, 'legend')
                    for axis_name in ('category_axis', 'value_axis'):
                        try:
                            axis = getattr(chart, axis_name)
                        except (ValueError, AttributeError):
                            continue  # e.g. a pie chart has no axes
                        if axis.visible:
                            chart_font(axis.tick_labels.font, s.shape_id, axis_name)
                        if axis.has_title:
                            type_check(axis.axis_title.text_frame, s.shape_id, axis_name + ' title')
                    for plot_number, plot in enumerate(chart.plots, 1):
                        if plot.has_data_labels:
                            chart_font(plot.data_labels.font, s.shape_id, f'plot {plot_number} labels')
            if s.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pw, ph = s.image.size
                crop = [s.crop_left, s.crop_top, s.crop_right, s.crop_bottom]
                ppi = min(pw / (s.width / EMU_IN), ph / (s.height / EMU_IN)) if s.width > 0 and s.height > 0 else 0
                # Compare the visible source aspect after any crop, not the whole raster.
                vw, vh = pw * (1-crop[0]-crop[2]), ph * (1-crop[1]-crop[3])
                stretch = abs((s.width / s.height) / (vw / vh) - 1) if s.height > 0 and vw > 0 and vh > 0 else 1
                image = {'shape': s.shape_id, 'name': s.name, 'sha256': hashlib.sha256(s.image.blob).hexdigest(),
                         'pixels': [pw, ph], 'crop': crop, 'approx_ppi': round(ppi, 1),
                         'source_url': s.click_action.hyperlink.address, 'inside_group': grouped}
                entry['pictures'].append(image)
                if any(crop):
                    warnings.append({'slide': i, 'shape': s.shape_id, 'kind': 'cropped_image_review'})
                if stretch > .02:
                    warnings.append({'slide': i, 'shape': s.shape_id, 'kind': 'stretched_image', 'relative_difference': round(stretch, 3)})
                if not grouped and ppi < 100:
                    warnings.append({'slide': i, 'shape': s.shape_id, 'kind': 'low_raster_resolution', 'approx_ppi': round(ppi, 1)})
        for shape in slide.shapes:
            inspect(shape)
        for a, (aid, ab) in enumerate(text_regions):
            for bid, bb in text_regions[a+1:]:
                width = min(ab[2], bb[2]) - max(ab[0], bb[0])
                height = min(ab[3], bb[3]) - max(ab[1], bb[1])
                if width > 1 and height > 1:
                    warnings.append({'slide': i, 'shapes': [aid, bid], 'kind': 'intersecting_text_regions',
                                     'intersection_pt': [round(width, 2), round(height, 2)]})
        if slide.has_notes_slide:
            entry['notes'] = slide.notes_slide.notes_text_frame.text
        if spec is not None and i <= len(spec.get('slides', [])):
            planned = slide_asset_ids(spec['slides'][i-1])
            expected = [load_asset(spec, key, base) for key in planned]
            if Counter(a['sha256'] for a in expected) != Counter(a['sha256'] for a in entry['pictures']):
                errors.append(f'Slide {i}: planned/embedded image bytes or occurrence counts differ')
            for key, asset in zip(planned, expected):
                candidates = [a for a in entry['pictures'] if a['sha256'] == asset['sha256']]
                if any(any(a['crop']) for a in candidates):
                    errors.append(f'Slide {i}: contain-only starter asset {key} has been cropped')
                for field in ('credit', 'license'):
                    if asset[field] not in entry['notes']:
                        errors.append(f'Slide {i}: image {key} lacks {field} in notes')
                if asset.get('source_url'):
                    if asset['source_url'] not in entry['notes'] or any(a['source_url'] != asset['source_url'] for a in candidates):
                        errors.append(f'Slide {i}: image {key} source link missing')
            entry['planned_asset_ids'] = planned
        slides.append(entry)
    if not slides:
        errors.append('No slides')
    return {'sha256': hashlib.sha256(payload).hexdigest(), 'slide_count': len(slides),
            'status': 'failed' if errors else ('review_required' if warnings else 'passed'),
            'errors': errors, 'warnings': warnings, 'slides': slides,
            'review_policy': {'min_font_pt': min_font_pt, 'protected_regions': regions},
            'scope': 'Structural/plan checks and heuristic risks only. Native review is required. Font sizes are explicit values, not all effective sizes. Raster text and inherited master/layout elements are not scanned; protected regions must come from template inspection. Group coordinates are not transformed. No-warning status is not proof of audience readability or submission readiness.'}


if __name__ == '__main__':
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('file', type=Path)
    p.add_argument('--spec', type=Path, help='Optional bounded generator JSON for image placement/provenance parity')
    p.add_argument('--min-font-pt', type=float, help='Flag explicit type below this threshold; e.g. 22 for projection')
    p.add_argument('--protected-regions', type=Path, help='JSON array of template regions in inches requiring overlap review')
    a = p.parse_args()
    try:
        spec = json.loads(a.spec.read_text(encoding='utf-8')) if a.spec else None
        regions = json.loads(a.protected_regions.read_text(encoding='utf-8')) if a.protected_regions else None
        result = check(a.file, spec, a.spec.parent if a.spec else None, a.min_font_pt, regions)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        raise SystemExit(1 if result['errors'] else 0)
    except (OSError, ValueError) as e:
        p.exit(1, str(e) + '\n')
