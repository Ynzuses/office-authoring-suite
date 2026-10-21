#!/usr/bin/env python3
"""Create a new editable PPTX from a bounded JSON layout contract."""
import argparse
import io
import json
import math
from pathlib import Path
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from lxml import etree
from pptx_assets import load_asset, slide_asset_ids


def keys(obj, allowed):
    unknown = set(obj) - set(allowed.split())
    if unknown:
        raise ValueError(f"Unsupported fields: {sorted(unknown)}")


def text(slide, value, x, y, w, h, size, color, font, bold=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(value.split("\n")):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = line
        para.font.name = font
        para.font.size = Pt(size)
        para.font.bold = bold
        para.font.color.rgb = RGBColor.from_string(color)
        para.space_after = Pt(0)
        para.space_before = Pt(0)
        para.line_spacing = 1.05
    return shape


def rect(slide, x, y, w, h, color, shape_type=MSO_SHAPE.RECTANGLE):
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    shape._element.spPr.append(OxmlElement('a:effectLst'))
    return shape


def picture(slide, asset_id, asset, x, y, w, h):
    pw, ph = asset['pixels']
    scale = min(w / pw, h / ph)
    iw, ih = pw * scale, ph * scale
    shape = slide.shapes.add_picture(str(asset['resolved_path']),
        Inches(x + (w-iw)/2), Inches(y + (h-ih)/2), width=Inches(iw), height=Inches(ih))
    shape.name = 'asset:' + asset_id
    shape._element.nvPicPr.cNvPr.set('descr', asset.get('alt', asset['credit']))
    if asset.get('source_url'):
        shape.click_action.hyperlink.address = asset['source_url']
    return shape


def plain_hyperlinks(prs):
    # Only applied to a newly created plain deck, never a supplied template.
    for part in prs.part.package.iter_parts():
        if str(part.partname).startswith('/ppt/theme/'):
            root = etree.fromstring(part.blob)
            ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
            for name in ('hlink', 'folHlink'):
                node = root.find('.//a:clrScheme/a:' + name, ns)
                for child in list(node):
                    node.remove(child)
                color = OxmlElement('a:srgbClr')
                color.set('val', '000000')
                node.append(color)
            part._blob = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)


def build(spec, base=Path('.')):
    keys(spec, "title font accent style slides assets")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.title = spec.get("title", "")
    prs.core_properties.author = ""
    prs.core_properties.last_modified_by = ""
    font, accent = spec.get("font", "Arial"), spec.get("accent", "117C83")
    RGBColor.from_string(accent)  # Reject invalid color before writing a file.
    dark, muted, pale = "213D4D", "50616B", "EDF4F4"
    style = spec.get('style', 'accent')
    if style not in ('accent', 'plain'):
        raise ValueError('style must be accent or plain')
    if style == 'plain':
        dark = muted = accent = '000000'
        pale = 'FFFFFF'
        plain_hyperlinks(prs)
    slides = spec.get("slides")
    if not isinstance(slides, list) or not slides:
        raise ValueError("slides must be a nonempty list")
    for number, item in enumerate(slides, 1):
        kind = item.get("layout")
        allowed = "layout title source notes " + {
            "cover": "body", "content": "body", "comparison": "items",
            "process": "items", "chart": "categories values series unit",
            "image": "image body caption side"
        }.get(kind, "")
        keys(item, allowed)
        if kind not in ("cover", "content", "comparison", "process", "chart", "image"):
            raise ValueError(f"Unsupported layout: {kind}")
        asset_ids = slide_asset_ids(item)
        assets = {key: load_asset(spec, key, base) for key in asset_ids}
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
        if style != 'plain':
            rect(slide, 0.6, 0.55, 0.7, 0.09, accent)
        text(slide, item["title"], 0.6, 0.9 if kind != "cover" else 1.65,
             12.1, 1.2 if kind == "cover" else 1.05, 36 if kind == "cover" else 30, dark, font, True)
        if kind in ("cover", "content"):
            text(slide, item["body"], 0.65, 3.3 if kind == "cover" else 2.25,
                 11.7, 2.6 if kind == "cover" else 3.9, 24, muted, font)
        elif kind == 'image':
            side = item.get('side', 'left')
            if side not in ('left', 'right'):
                raise ValueError('Image side must be left or right')
            image_x, body_x = (0.6, 8.25) if side == 'left' else (5.55, 0.6)
            picture(slide, item['image'], assets[item['image']], image_x, 2.2, 7.15, 3.2)
            text(slide, item.get('caption', ''), image_x, 5.6, 7.15, .72, 22, dark, font)
            text(slide, item['body'], body_x, 2.28, 4.45, 3.93, 24, dark, font)
        elif kind in ("comparison", "process"):
            items = item["items"]
            if (kind == "comparison" and len(items) != 2) or (kind == "process" and not 2 <= len(items) <= 4):
                raise ValueError("comparison requires 2 items; process requires 2–4")
            n = len(items)
            gap = 0.45
            width = (12.1 - gap * (n - 1)) / n
            with_images = bool(asset_ids)
            for i, cell in enumerate(items):
                keys(cell, "title body image")
                x = 0.6 + i * (width + gap)
                rect(slide, x, 2.2, width, 3.75, pale)
                if cell.get('image'):
                    picture(slide, cell['image'], assets[cell['image']], x+.22, 2.38, width-.44, 1.2)
                text(slide, cell["title"], x + 0.22, 3.78 if with_images else 2.52,
                     width - 0.44, 0.72, 24, accent, font, True)
                text(slide, cell["body"], x + 0.22, 4.62 if with_images else 3.43,
                     width - 0.44, 1.18 if with_images else 2.05, 22, dark, font)
                if kind == "process" and i < n - 1:
                    rect(slide, x + width + 0.07, 3.65, 0.3, 0.25, accent, MSO_SHAPE.CHEVRON)
        else:
            categories, values = item["categories"], item["values"]
            if not categories or len(categories) != len(values) or any(not isinstance(v, (float, int)) or not math.isfinite(v) for v in values):
                raise ValueError("Chart requires matching categories and finite numeric values")
            data = CategoryChartData()
            data.categories = categories
            data.add_series(item["series"], values)
            chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,
                Inches(0.85), Inches(2.3), Inches(11.6), Inches(3.85), data).chart
            chart.has_legend = False
            chart.chart_style = 10
            chart.value_axis.has_title = True
            chart.value_axis.axis_title.text_frame.text = item["unit"]
            for para in chart.value_axis.axis_title.text_frame.paragraphs:
                para.font.size = Pt(22)
                para.font.name = font
            chart.value_axis.tick_labels.font.size = Pt(22)
            chart.category_axis.tick_labels.font.size = Pt(22)
            plot = chart.plots[0]
            plot.has_data_labels = True
            plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
            plot.data_labels.font.size = Pt(22)
            series = chart.series[0]
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = RGBColor.from_string(accent)
        text(slide, item.get("source", ""), 0.6, 6.46 if assets else 6.8, 11.4, .32, 12, muted, font)
        text(slide, str(number), 12.15, 6.8, 0.55, 0.4, 12, muted, font)
        notes = item.get('notes', '')
        if assets:
            credit = 'Images: ' + ' | '.join(a['credit'] + '; ' + a['license'] for a in assets.values())
            text(slide, credit, .6, 6.87, 11.4, .45, 12, muted, font)
            for key, asset in assets.items():
                notes += '\n\nImage ' + key + ': ' + asset['credit'] + '\nLicense: ' + asset['license']
                if asset.get('source_url'):
                    notes += '\nSource: ' + asset['source_url']
                notes += '\nEmbedded source SHA-256: ' + asset['sha256']
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    return prs


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("spec", type=Path)
    p.add_argument("output", type=Path)
    a = p.parse_args()
    try:
        if a.output.suffix.lower() != ".pptx":
            raise ValueError("Output must end in .pptx")
        prs = build(json.loads(a.spec.read_text(encoding="utf-8")), a.spec.parent)
        buffer = io.BytesIO()
        prs.save(buffer)
        a.output.parent.mkdir(parents=True, exist_ok=True)
        with a.output.open("xb") as f:
            f.write(buffer.getvalue())
        print(a.output)
    except (OSError, ValueError, KeyError, TypeError) as e:
        p.exit(1, str(e) + "\n")


if __name__ == "__main__":
    main()
