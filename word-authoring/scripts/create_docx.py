#!/usr/bin/env python3
"""Create a new DOCX from the documented JSON subset. Never overwrite input."""
import argparse
import io
import json
from pathlib import Path
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def keys(obj, allowed):
    unknown = set(obj) - set(allowed.split())
    if unknown:
        raise ValueError(f"Unsupported fields: {sorted(unknown)}")


def fill_paragraph(p, block):
    if "text" in block and "runs" in block:
        raise ValueError("Choose text or runs, not both")
    for item in block.get("runs", [{"text": block.get("text", "")} ]):
        keys(item, "text bold italic superscript subscript")
        run = p.add_run(item["text"])
        for name in ("bold", "italic", "superscript", "subscript"):
            if name in item:
                setattr(run.font, name, item[name])


def build(spec, base):
    keys(spec, "title subtitle language font east_asian_font body_pt page_size margin_cm blocks")
    doc = Document()
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.title = spec.get("title", "")
    doc.core_properties.language = spec.get("language", "en")
    section = doc.sections[0]
    page_size = spec.get("page_size", "A4")
    if page_size not in ("A4", "Letter"):
        raise ValueError("page_size must be A4 or Letter")
    w, h = (21, 29.7) if page_size == "A4" else (21.59, 27.94)
    margin = float(spec.get("margin_cm", 2.3))
    if not 0.5 <= margin < min(w, h) / 3:
        raise ValueError("margin_cm would leave inadequate page space")
    section.page_width, section.page_height = Cm(w), Cm(h)
    section.top_margin = section.bottom_margin = Cm(margin)
    section.left_margin = section.right_margin = Cm(margin)
    font = spec.get("font", "Arial")
    ea_font = spec.get("east_asian_font", "Arial Unicode MS" if sys.platform == "darwin" else "Noto Sans CJK SC")
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Number", "Caption"):
        style = doc.styles[name]
        style.font.name = font
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), ea_font)
    # The blank template includes theme font attributes that override explicit
    # names in Word. Normalize this new document's style/default font mappings.
    # This is not a safe generic operation on an existing rich document.
    for fonts in doc.styles.element.iter(qn("w:rFonts")):
        for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn("w:" + attribute), None)
        for attribute, value in (("ascii", font), ("hAnsi", font), ("eastAsia", ea_font)):
            fonts.set(qn("w:" + attribute), value)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(float(spec.get("body_pt", 11)))
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.15
    for level, size in ((1, 16), (2, 13), (3, 11)):
        style = doc.styles[f"Heading {level}"]
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string("213D4D")
        style.paragraph_format.keep_with_next = True
    if spec.get("title"):
        doc.add_heading(spec["title"], 0)
    if spec.get("subtitle"):
        doc.add_paragraph(spec["subtitle"], "Subtitle")
    blocks = spec.get("blocks")
    if not isinstance(blocks, list) or not blocks:
        raise ValueError("blocks must be a nonempty list")
    for b in blocks:
        kind = b.get("type")
        if kind in ("paragraph", "heading"):
            keys(b, "type text runs level")
            level = int(b.get("level", 1))
            if kind == "heading" and level not in (1, 2, 3):
                raise ValueError("Starter supports heading levels 1–3")
            p = doc.add_heading(level=level) if kind == "heading" else doc.add_paragraph()
            fill_paragraph(p, b)
        elif kind == "list":
            keys(b, "type items ordered")
            for item in b["items"]:
                doc.add_paragraph(item, "List Number" if b.get("ordered") else "List Bullet")
        elif kind == "table":
            keys(b, "type headers rows")
            headers = b["headers"]
            if not headers or any(len(r) != len(headers) for r in b["rows"]):
                raise ValueError("Table rows must match the nonempty headers")
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = "Light Shading Accent 1"
            for cell, value in zip(table.rows[0].cells, headers):
                cell.text = str(value)
            repeat = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for row in b["rows"]:
                for cell, value in zip(table.add_row().cells, row):
                    cell.text = str(value)
        elif kind == "image":
            keys(b, "type path width_cm caption")
            image = (base / b["path"]).resolve(strict=True)
            width = float(b.get("width_cm", w - 2 * margin))
            if not 0 < width <= w - 2 * margin:
                raise ValueError("Image width exceeds text area")
            doc.add_picture(str(image), width=Cm(width))
            if b.get("caption"):
                doc.add_paragraph(b["caption"], "Caption")
        elif kind == "page_break":
            keys(b, "type")
            doc.add_page_break()
        else:
            raise ValueError(f"Unsupported block type: {kind}")
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    # Use explicit run mappings for this new document. Word can otherwise use
    # an inherited East Asian font even when the paragraph style names one.
    for part in (doc.element, section.footer._element):
        for run in part.iter(qn("w:r")):
            fonts = run.get_or_add_rPr().get_or_add_rFonts()
            fonts.set(qn("w:ascii"), font)
            fonts.set(qn("w:hAnsi"), font)
            fonts.set(qn("w:eastAsia"), ea_font)
    return doc


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("spec", type=Path)
    p.add_argument("output", type=Path)
    a = p.parse_args()
    try:
        if a.output.suffix.lower() != ".docx":
            raise ValueError("Output must end in .docx")
        doc = build(json.loads(a.spec.read_text(encoding="utf-8")), a.spec.resolve().parent)
        buf = io.BytesIO()
        doc.save(buf)
        a.output.parent.mkdir(parents=True, exist_ok=True)
        with a.output.open("xb") as stream:
            stream.write(buf.getvalue())
        print(a.output)
    except (OSError, ValueError, KeyError, TypeError) as e:
        p.exit(1, str(e) + "\n")


if __name__ == "__main__":
    main()
